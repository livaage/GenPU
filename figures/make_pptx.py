"""Build an editable PowerPoint of the model figure.

The graphic goes in as a high-resolution render of `model_arch_notext.svg` (which contains NO
text), and every label from `model_arch_text.json` is re-created as a NATIVE PowerPoint text box at
the same position — so all text is selectable, editable and restyleable in PowerPoint/Keynote/
Google Slides, while the diagram itself stays crisp.
"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

HERE = Path(__file__).resolve().parent
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5                  # 16:9
FONT = "Arial"

d = json.loads((HERE / "model_arch_text.json").read_text())
W, H, items = d["w"], d["h"], d["items"]

# fit the figure inside the slide, preserving aspect
img_aspect, slide_aspect = W / H, SLIDE_W_IN / SLIDE_H_IN
if img_aspect >= slide_aspect:                        # width-limited
    img_w_in, img_h_in = SLIDE_W_IN, SLIDE_W_IN / img_aspect
else:                                                 # height-limited
    img_h_in, img_w_in = SLIDE_H_IN, SLIDE_H_IN * img_aspect
off_x_in, off_y_in = (SLIDE_W_IN - img_w_in) / 2, (SLIDE_H_IN - img_h_in) / 2
PPT_PER_PX = img_w_in * 72.0 / W                      # points per SVG pixel

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SLIDE_W_IN), Inches(SLIDE_H_IN)
blank = prs.slide_layouts[6]

slide = prs.slides.add_slide(blank)
slide.shapes.add_picture(str(HERE / "model_arch_notext@4x.png"),
                         Inches(off_x_in), Inches(off_y_in),
                         width=Inches(img_w_in), height=Inches(img_h_in))

def place(it):
    size_pt = it["size"] * PPT_PER_PX
    x_pt = off_x_in * 72.0 + it["x"] * PPT_PER_PX
    y_pt = off_y_in * 72.0 + it["y"] * PPT_PER_PX
    w_pt = max(len(it["text"]) * size_pt * 0.62, 24.0) + 12.0
    left = {"start": x_pt, "end": x_pt - w_pt, "middle": x_pt - w_pt / 2}[it["anchor"]]
    align = {"start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT, "middle": PP_ALIGN.CENTER}[it["anchor"]]
    # SVG y is the BASELINE; a PowerPoint box is positioned by its top edge
    top = y_pt - size_pt * 1.02
    box = slide.shapes.add_textbox(Pt(left), Pt(top), Pt(w_pt), Pt(size_pt * 1.55))
    tf = box.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = it["text"]
    f = r.font
    f.name, f.size, f.bold = FONT, Pt(round(size_pt, 1)), it["bold"]
    f.color.rgb = RGBColor.from_string(it["fill"].lstrip("#").upper())

for it in items:
    place(it)

# a reference slide with the original labelled render, in case an edit goes wrong
ref = prs.slides.add_slide(blank)
ref.shapes.add_picture(str(HERE / "model_arch.png"), Inches(off_x_in), Inches(off_y_in),
                       width=Inches(img_w_in), height=Inches(img_h_in))

out = HERE / "model_arch.pptx"
prs.save(str(out))
print(f"wrote {out}")
print(f"  figure {W}x{H}px -> {img_w_in:.2f}x{img_h_in:.2f}in at ({off_x_in:.2f},{off_y_in:.2f})in")
print(f"  {len(items)} editable text boxes, {PPT_PER_PX:.4f} pt/px "
      f"(title {max(i['size'] for i in items)*PPT_PER_PX:.1f}pt, "
      f"smallest {min(i['size'] for i in items)*PPT_PER_PX:.1f}pt)")
